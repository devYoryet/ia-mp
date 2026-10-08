"""Panel de revisión con una BD falsa: login, protección, escape, decisiones, Excel."""
import io
import os
from datetime import datetime

import pytest

os.environ.setdefault("F2_SESSION_SECRET", "x" * 32)
bcrypt = pytest.importorskip("bcrypt")
pytest.importorskip("fastapi")
from fastapi.testclient import TestClient  # noqa: E402

import panel  # noqa: E402
from bd import validar_escritura  # noqa: E402

CLAVE = "secreta123"
HASH_LARAVEL = "$2y$" + bcrypt.hashpw(CLAVE.encode(), bcrypt.gensalt(rounds=4)).decode()[4:]
FILA = {
    "id": 11, "tabla_origen": "compra_agil", "fila_id": 999, "licitacion": "1234-56-COT26",
    "fecha_publicacion": datetime(2026, 10, 1, 10, 0), "fecha_cierre": datetime(2026, 10, 20, 15, 0),
    "descripcion": "GUANTE NITRILO <script>alert(1)</script>\x0b CAJA 100", "titulo": "Insumos",
    "codigo_onu": "42132203", "nombre_onu": "Guantes médicos de examen", "ia_interes": 0,
    "ia_metodo": "descarte_item", "estado_gestor": 0, "clasificador_f1": "Carolina Burgos",
    "categoria": "DEV-GUA", "categoria_nombre": "Guantes médicos", "linea": "Device", "subcategoria": "Guante",
    "terminos": "Guante", "senal": "ambas", "estado_auto": "verde", "otras_categorias": None, "vigente": 1,
    "motivo_no_vigente": None, "decision": None, "categoria_final": None, "revisado_por": None,
    "revisado_en": None, "motivo": None, "cerrada": 0,
}
CATS = [{"codigo": "SRV-MAN", "linea": "Servicio técnico", "nombre": "Mantención de equipos médicos"},
        {"codigo": "DEV-GUA", "linea": "Device", "nombre": "Guantes médicos"}]


class FakeBD:
    def __init__(self):
        self.escrituras = []

    def todos(self, sql, params=()):
        s = " ".join(sql.split())
        if "FROM users" in s:
            return [{"id": 5, "name": "Carolina Burgos", "email": "c@x.cl", "password": HASH_LARAVEL}] \
                if params[0] == "c@x.cl" else []
        if "FROM clasificador_f2_categorias" in s:
            return CATS
        if "FROM clasificador_f2_corridas" in s:
            return [{"tipo": "rapido", "creado_en": datetime(2026, 10, 8, 16, 47), "segundos": 9, "error": None,
                     "hace": 3, "dias": 3, "filas_leidas": 1, "calzan": 1, "insertadas": 0, "actualizadas": 0,
                     "anuladas": 0}]
        if "GROUP BY 1,2,3" in s:  # precisión medida
            return [{"categoria": "DEV-GUA", "estado_auto": "verde", "senal": "ambas", "a": 9, "r": 1}]
        if "GROUP BY 1" in s and "AS c" in s:
            return [{"c": "DEV-GUA", "n": 1, "verde": 1, "revision": 0, "aprob": 0, "rech": 0}]
        if "tabla_origen AS t" in s:
            return [{"t": "compra_agil", "verde": 1, "revision": 0}]
        if "COUNT(*) AS n FROM clasificador_f2_resultado" in s:
            return [{"n": 1}]
        if "SUM(vigente=1 AND decision IS NULL" in s:
            return [{"revision": 3, "verde": 1, "revisadas": 0, "anuladas": 0, "hoy": 0}]
        if "GROUP BY 1" in s and "AS c" in s:
            return [{"c": "DEV-GUA", "n": 1, "verde": 1, "revision": 0, "aprob": 0, "rech": 0}]
        if "WHERE id=%s" in s:
            return [{"id": 11, "categoria": "DEV-GUA", "vigente": 1}] if params[0] == 11 else []
        if "FROM clasificador_f2_resultado" in s:
            return [FILA]
        return []

    def uno(self, sql, params=()):
        r = self.todos(sql, params)
        return r[0] if r else None

    def transaccion(self, sentencias):
        for sql, p in sentencias:
            validar_escritura(sql)  # el mismo candado que la BD real
            self.escrituras.append((sql, p))

    def cerrar(self):
        pass


@pytest.fixture
def cli(monkeypatch):
    fake = FakeBD()
    monkeypatch.setattr(panel, "get_bd", lambda: fake)
    c = TestClient(panel.app)
    c.fake = fake
    return c


def _login(c):
    r = c.post("/login", data={"email": "c@x.cl", "password": CLAVE, "next": "/revision"}, follow_redirects=False)
    assert r.status_code == 303


def test_salud_es_publica(cli):
    assert cli.get("/salud").json()["ok"] is True


def test_sin_sesion_redirige_a_login(cli):
    r = cli.get("/revision?vista=verde", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/login")
    assert cli.post("/decidir", data={"id": 11, "decision": "aprobado"}, follow_redirects=False).status_code == 303


def test_login_rechaza_clave_mala(cli):
    r = cli.post("/login", data={"email": "c@x.cl", "password": "otra"}, follow_redirects=False)
    assert r.status_code == 401


def test_revision_escapa_html(cli):
    _login(cli)
    r = cli.get("/revision?vista=verde")
    assert r.status_code == 200
    assert "<script>alert(1)</script>" not in r.text and "&lt;script&gt;" in r.text
    assert "Guantes médicos" in r.text and "Aprobar" in r.text


def test_decidir_escribe_solo_columnas_de_revision(cli):
    _login(cli)
    r = cli.post("/decidir", data={"id": 11, "decision": "rechazado", "categoria": "SRV-MAN", "motivo": "es de trabajo"},
                 headers={"Accept": "application/json"})
    assert r.json() == {"ok": True, "decision": "rechazado", "por": "Carolina Burgos"}
    sql, p = cli.fake.escrituras[-1]
    assert sql.startswith("UPDATE clasificador_f2_resultado SET decision=%s, categoria_final=%s, revisado_por=%s")
    assert p == ("rechazado", "SRV-MAN", "Carolina Burgos", "es de trabajo", 11)


def test_decidir_valida(cli):
    _login(cli)
    assert cli.post("/decidir", data={"id": 11, "decision": "borrar"}, headers={"Accept": "application/json"}).status_code == 400
    assert cli.post("/decidir", data={"id": 11, "decision": "aprobado", "categoria": "XXX"},
                    headers={"Accept": "application/json"}).status_code == 400
    assert cli.post("/decidir", data={"id": 77, "decision": "aprobado"}, headers={"Accept": "application/json"}).status_code == 404
    assert cli.fake.escrituras == []


def test_misma_categoria_no_se_guarda_como_correccion(cli):
    _login(cli)
    cli.post("/decidir", data={"id": 11, "decision": "aprobado", "categoria": "DEV-GUA"}, headers={"Accept": "application/json"})
    assert cli.fake.escrituras[-1][1][1] is None


def test_exportar_excel(cli):
    from openpyxl import load_workbook
    _login(cli)
    r = cli.get("/exportar.xlsx?vista=verde")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["A1"].value == "Fuente" and ws["F2"].value == "Guantes médicos" and "\x0b" not in ws["M2"].value


def test_resumen(cli):
    _login(cli)
    r = cli.get("/resumen")
    assert r.status_code == 200 and "Precisión medida" in r.text and "90%" in r.text


def test_next_no_permite_redireccion_externa():
    assert panel._next_seguro("//evil.com") == "/revision"
    assert panel._next_seguro("https://evil.com") == "/revision"
    assert panel._next_seguro("/revision?vista=verde") == "/revision?vista=verde"
